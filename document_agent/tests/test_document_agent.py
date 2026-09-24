import json
import subprocess
import sys
from pathlib import Path
import pytest
from document_agent.document_loader import Document, load_document
from document_agent.agent import run_agent
from document_agent.tools import DocumentTools


def make_pdf(path):
    objects=[b'<< /Type /Catalog /Pages 2 0 R >>',b'<< /Type /Pages /Kids [3 0 R 5 0 R] /Count 2 >>']
    for page, text in [(3,'Alice must submit report by Friday.'),(5,'Bob must review report by Monday.')]:
        objects.append(f'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 7 0 R >> >> /Contents {page+1} 0 R >>'.encode())
        stream=f'BT /F1 12 Tf 40 770 Td (Project Header) Tj 0 -70 Td ({text}) Tj ET'.encode()
        objects.append(b'<< /Length '+str(len(stream)).encode()+b' >>\nstream\n'+stream+b'\nendstream')
    objects.append(b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>')
    data=b'%PDF-1.4\n';offsets=[0]
    for i,obj in enumerate(objects,1):
        offsets.append(len(data));data+=f'{i} 0 obj\n'.encode()+obj+b'\nendobj\n'
    xref=len(data);data+=b'xref\n0 8\n0000000000 65535 f \n'
    data+=b''.join(f'{offset:010d} 00000 n \n'.encode() for offset in offsets[1:])
    data+=f'trailer << /Size 8 /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF'.encode()
    path.write_bytes(data)


def test_pdf_multipage_headers_and_location(tmp_path):
    p=tmp_path/'sample.pdf';make_pdf(p)
    doc=load_document(p)
    assert len(doc.chunks)==2
    assert doc.chunks[1]['location']=='page 2'
    assert 'Bob must' in doc.chunks[1]['text']
    assert all('Project Header' not in c['text'] for c in doc.chunks)


def test_word_body_order_tables_and_ignored_headers(tmp_path):
    from docx import Document as Word
    d=Word();d.sections[0].header.paragraphs[0].text='Header'
    d.add_paragraph('First task')
    t=d.add_table(rows=1,cols=2);t.cell(0,0).text='Owner';t.cell(0,1).text='Alice'
    d.add_paragraph('Last task')
    p=tmp_path/'test.docx';d.save(p)
    doc=load_document(p)
    assert [c['text'] for c in doc.chunks]==['First task','Owner | Alice','Last task']


@pytest.mark.parametrize('suffix',['.pdf','.docx','.txt'])
def test_invalid_file_fails_cleanly(tmp_path,suffix):
    p=tmp_path/('bad'+suffix);p.write_text('not a document')
    with pytest.raises(ValueError):load_document(p)


class Scripted:
    def __init__(self, actions, mode):self.actions=iter(actions);self.mode=mode;self.messages=[]
    def __call__(self,messages):
        self.messages=list(messages)
        args=next(self.actions)
        if self.mode=='json':return {'content':json.dumps(args)}
        return {'tool_calls':[{'id':'call'+str(len(messages)),'type':'function','function':{'name':'document_action','arguments':json.dumps(args)}}]}


@pytest.mark.parametrize('mode',['native','json'])
def test_loop_observes_error_retries_and_completes(mode):
    doc=Document('test',[{'id':1,'location':'page 1','text':'Alice must send report by Friday.'}],[])
    actions=[{'action':'finish','reason':'Attempt completion'},
             {'action':'read_chunk','chunk_id':1,'reason':'Review remaining evidence'},
             {'action':'record_actions','reason':'Save task','items':[{'chunk_id':1,'quote':doc.chunks[0]['text'],'owner':'Alice','due_date':'Friday'}]},
             {'action':'finish','reason':'All evidence reviewed'}]
    client=Scripted(actions,mode);events=[]
    result=run_agent(doc,'extract_action_items',client,log=events.append)
    assert result['status']=='awaiting_review' and len(result['items'])==1
    assert 'error' in events[0]['observation']
    assert any('observation' in (m.get('content') or '').lower() or m['role']=='tool' for m in client.messages)


def test_quotes_and_fields_are_checked_atomically():
    doc=Document('test',[{'id':1,'location':'block 1','text':'Alice must send report.'}],[])
    t=DocumentTools(doc)
    with pytest.raises(ValueError):t.execute({'action':'finish','reason':'Done'})
    t.execute({'action':'read_chunk','chunk_id':1,'reason':'Read'})
    for quote,owner in [('Invented','Alice'),('Alice must send report.','Bob')]:
        with pytest.raises(ValueError):t.execute({'action':'record_actions','reason':'Save','items':[{'chunk_id':1,'quote':quote,'owner':owner,'due_date':None}]})
    assert t.items==[]
    r=run_agent(doc,'extract_requirements',Scripted([{'action':'read_chunk','chunk_id':1,'reason':'Read'}],'json'),max_steps=1)
    assert r['status']=='step_limit'


def test_ingest_cli_and_output_no_overwrite(tmp_path):
    from docx import Document as Word
    d=Word();d.add_paragraph('Review the report.');p=tmp_path/'test.docx';d.save(p)
    out=tmp_path/'output'
    command=[sys.executable,'-m','document_agent.main','--file',str(p),'--ingest-only','--output',str(out)]
    assert subprocess.run(command,capture_output=True).returncode==0
    assert json.loads((out/'document.json').read_text())['chunks']
    assert subprocess.run(command,capture_output=True).returncode==1


@pytest.mark.parametrize('decision',['approve','reject'])
def test_durable_review_survives_new_process_without_model(tmp_path,decision):
    from langgraph.checkpoint.sqlite import SqliteSaver
    from document_agent.agent import build_agent_graph, initial_state, result_from_state
    from document_agent.main import save_artifacts
    doc=Document('test',[{'id':1,'location':'page 1','text':'Alice must send report.'}],[])
    client=Scripted([{'action':'read_chunk','chunk_id':1,'reason':'Read'},
                     {'action':'record_actions','reason':'Draft','items':[{'chunk_id':1,'quote':doc.chunks[0]['text'],'owner':'Alice','due_date':None}]},
                     {'action':'finish','reason':'Review complete'}],'native')
    out=tmp_path/'run';out.mkdir()
    (out/'run.json').write_text(json.dumps({'thread_id':'test','max_steps':10}))
    with SqliteSaver.from_conn_string(str(out/'checkpoint.sqlite')) as saver:
        graph=build_agent_graph(client,saver)
        state=graph.invoke(initial_state(doc,'extract_action_items','native',10),{'configurable':{'thread_id':'test'}})
        assert state['__interrupt__']
        save_artifacts(out,result_from_state(state))
    assert not (out/'result.json').exists()
    command=[sys.executable,'-m','document_agent.main','--resume',str(out),'--decision',decision]
    completed=subprocess.run(command,capture_output=True,text=True)
    assert completed.returncode==0,completed.stderr
    review=json.loads((out/'review.json').read_text())
    assert review['status']==('approved' if decision=='approve' else 'rejected')
    assert (out/'result.json').exists()==(decision=='approve')
    assert (out/'report.md').exists()==(decision=='approve')
    assert subprocess.run(command,capture_output=True).returncode==1


def test_model_error_is_terminal_and_preserves_drafts():
    def fail(messages):raise RuntimeError('Model unavailable')
    fail.mode='json'
    doc=Document('test',[{'id':1,'location':'page 1','text':'Read this.'}],[])
    result=run_agent(doc,'extract_requirements',fail)
    assert result['status']=='model_error'
    assert result['items']==[] and len(result['steps'])==1
